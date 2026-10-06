from pathlib import Path
import collections
import hashlib
import json
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from scipy.stats import rankdata
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent if (ROOT.parent / 'src/seq2lead').exists() else ROOT.parent / 'closeout/seq2lead'
RUN = REPO / 'data/asof/m11h/run-20261005T134842Z'
FIG = REPO / 'paper/figures'
FIG.mkdir(exist_ok=True)
res = json.loads((RUN / 'results.json').read_text())
manifest = json.loads((RUN / 'manifest.json').read_text())
for key, filename in [('predictions', 'predictions.npz'), ('evaluation_table', 'evaluation-pairs.jsonl')]:
    assert hashlib.sha256((RUN / filename).read_bytes()).hexdigest() == manifest[key]['sha256']
blob = np.load(RUN / 'predictions.npz', allow_pickle=True)
rows = [json.loads(x) for x in (RUN / 'evaluation-pairs.jsonl').read_text().splitlines()]
assert list(map(str, blob['pair'])) == [r['pair'] for r in rows]
families = ['B0-target-mean', 'B1-ligand-ecfp4-lgbm', 'B2-protein-esm2-lgbm', 'B3L-ligand-1nn', 'B4-concat-mlp', 'dual-encoder']
labels = ['Target mean', 'Ligand LightGBM', 'Protein LightGBM', 'Target-conditioned 1-NN', 'Concat MLP', 'Dual encoder']
colors = ['#8C9298', '#7295B5', '#ABB8C7', '#548B7F', '#174E77', '#C76939']
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False, 'svg.fonttype': 'none', 'figure.facecolor': 'white'})
provenance = {'inputs': {x: hashlib.sha256((RUN / x).read_bytes()).hexdigest() for x in ['predictions.npz', 'evaluation-pairs.jsonl', 'results.json']}, 'notes': ['Schematics are original explanatory diagrams, not molecular measurements.', 'Prediction plots reproduce saved scores without fitting or selection.', 'Seed points and ranges are not confidence intervals.', 'Per-target scatter is descriptive and does not test significance or equivalence.']}

def save(fig, name):
    fig.savefig(FIG / f'{name}.png', dpi=240, bbox_inches='tight', facecolor='white')
    fig.savefig(FIG / f'{name}.svg', bbox_inches='tight', facecolor='white')
    plt.close(fig)

def box(ax, x, y, w, h, text, color='#EAF0F5', size=11):
    p = FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.012,rounding_size=0.02', fc=color, ec='#647585', lw=1)
    ax.add_patch(p)
    ax.text(x + w / 2, y + h / 2, text, ha='center', va='center', fontsize=size)

def arrow(ax, a, b, color='#526575', style='-'):
    ax.add_patch(FancyArrowPatch(a, b, arrowstyle='-|>', mutation_scale=13, lw=1.35, color=color, linestyle=style))
fig, ax = plt.subplots(figsize=(10, 5.2))
ax.set(xlim=(0, 1), ylim=(0, 1))
ax.axis('off')
box(ax, 0.03, 0.75, 0.38, 0.2, 'Snapshot A · January 2026\nPinned archives → isolated SQL curation')
box(ax, 0.59, 0.75, 0.38, 0.2, 'Snapshot B · September 2026\nPreviously inspected local corpus')
box(ax, 0.29, 0.46, 0.42, 0.19, 'Counted matching within context slots\nContent keys + source row locators')
arrow(ax, (0.22, 0.75), (0.39, 0.65))
arrow(ax, (0.78, 0.75), (0.61, 0.65))
box(ax, 0.02, 0.16, 0.29, 0.21, 'Historical A only\nTrain / validation reservation\nNo future edits to training', size=10.5)
box(ax, 0.355, 0.16, 0.29, 0.21, 'Eligible added evidence\nIncrement forms labels\nCorrection candidates withheld', color='#F8EBDF', size=10.5)
box(ax, 0.69, 0.16, 0.29, 0.21, 'Complete actual B\nConsistency screen + audit\nCorrections remain in B', color='#E7F2ED', size=10.5)
arrow(ax, (0.3, 0.8), (0.165, 0.37))
arrow(ax, (0.47, 0.46), (0.5, 0.37))
arrow(ax, (0.78, 0.75), (0.835, 0.37))
ax.text(0.5, 0.055, 'Two increment interpretations × screened / unscreened branches = four reported cells', ha='center', fontsize=10)
save(fig, 'figure_1_evidence_workflow')
fig, ax = plt.subplots(figsize=(10, 5.1))
ax.set(xlim=(0, 1), ylim=(0, 1))
ax.axis('off')
box(ax, 0.02, 0.69, 0.27, 0.23, 'Compound structure\nChiral ECFP4 · 2,048 bits\nFixed representation', size=10.5)
box(ax, 0.02, 0.27, 0.27, 0.23, 'Protein sequence\nFrozen ESM-2 · 1,280 values\nMean over residues', size=10.5)
box(ax, 0.38, 0.69, 0.25, 0.23, 'Compound tower\nLearned projection\n512-dimensional zc', color='#E7F2ED')
box(ax, 0.38, 0.27, 0.25, 0.23, 'Protein tower\nA-train fitted transform\n512-dimensional zp', color='#E7F2ED')
arrow(ax, (0.29, 0.805), (0.38, 0.805))
arrow(ax, (0.29, 0.385), (0.38, 0.385))
box(ax, 0.73, 0.48, 0.25, 0.23, 'Affine cosine head\nŷ = a · cos(zc, zp) + b\nPredicted pKi', color='#F8EBDF')
arrow(ax, (0.63, 0.805), (0.76, 0.71))
arrow(ax, (0.63, 0.385), (0.76, 0.48))
ax.text(0.5, 0.11, 'Fit exact-Ki MSE on A-train → select by A-validation RMSE → score B once', ha='center', fontsize=11)
ax.text(0.5, 0.035, 'Independent towers permit precomputed library projections; ranking uses the affine score.', ha='center', fontsize=10)
save(fig, 'figure_2_dual_encoder')
primary = res['cells']['declared_increment/screened_primary']['groups']['new_to_fitting']['by_model']
values = []
fig, ax = plt.subplots(figsize=(9, 4.8))
for i, (fam, col) in enumerate(zip(families, colors)):
    vals = [v['metrics']['auroc'] for k, v in primary.items() if k == fam or k.startswith(fam + '-seed')]
    values.append(vals)
    y = 5 - i
    ax.plot([min(vals), max(vals)], [y, y], color=col, lw=3)
    ax.scatter(vals, [y] * len(vals), s=38, color=col, alpha=0.8, zorder=3)
    ax.scatter([np.mean(vals)], [y], s=85, marker='|', color='black', zorder=4)
    ax.text(0.825, y, f'{np.mean(vals):.4f}', va='center', fontsize=11)
ax.axvline(0.5, color='#969696', ls='--', lw=1)
ax.set(yticks=range(6), yticklabels=labels[::-1], xlim=(0.47, 0.88), xlabel='Macro AUROC on 123 rankable targets')
ax.grid(axis='x', alpha=0.15)
ax.set_title('Primary new-to-fitting cohort · 22,221 pairs across 992 targets', loc='left', fontsize=12, pad=15)
ax.text(0.0, -0.2, 'Dots: individual fits. Black tick: family mean. Lines: seed range, not a confidence interval.', transform=ax.transAxes, fontsize=9)
save(fig, 'figure_3_ranking_results')
new = {'new_absent_from_a', 'new_reserved_for_validation', 'new_a_present_excluded_from_fitting'}
bytarget = collections.defaultdict(list)
for i, r in enumerate(rows):
    arm = r['arms']['declared_increment']
    if r['stratum'] in new and arm['scoreable'] and arm['branches']['screened_primary']:
        bytarget[r['pair'].split('|')[1]].append(i)

def auc(y, s):
    n = int(y.sum())
    m = len(y) - n
    return (float(rankdata(s)[y].sum()) - n * (n + 1) / 2) / (n * m)
x = []
y = []
detail = []
for target, ix in sorted(bytarget.items()):
    iy = np.array([rows[i]['arms']['declared_increment']['label'] == 'active' for i in ix])
    if iy.sum() < 5 or (~iy).sum() < 5:
        continue
    vals = []
    for fam in ['B4-concat-mlp', 'dual-encoder']:
        vals.append([auc(iy, np.asarray(blob[k])[ix]) for k in blob.files if k.startswith(fam + '-seed')])
    x.append(np.mean(vals[0]))
    y.append(np.mean(vals[1]))
    detail.append({'target_sequence_sha256': target, 'pairs': len(ix), 'positives': int(iy.sum()), 'concat_seed_mean_auroc': x[-1], 'dual_seed_mean_auroc': y[-1]})
assert len(x) == 123
assert abs(np.mean(x) - np.mean(values[4])) < 1e-12 and abs(np.mean(y) - np.mean(values[5])) < 1e-12
(FIG / 'per_target_comparison.json').write_text(json.dumps(detail, indent=2) + '\n')
fig, axs = plt.subplots(1, 2, figsize=(10, 4.4), gridspec_kw={'width_ratios': [1, 1.25]})
ax = axs[0]
ax.plot([0, 1], [0, 1], ls='--', color='#999999', lw=1)
ax.scatter(x, y, c='#174E77', s=23, alpha=0.65)
ax.set(xlim=(0, 1), ylim=(0, 1), xlabel='Concat MLP AUROC', ylabel='Dual encoder AUROC')
ax.set_aspect('equal')
ax.set_title('a  Paired target comparison', loc='left', fontsize=11)
ax.text(0.03, 0.96, '123 targets · mean over seeds', transform=ax.transAxes, va='top', fontsize=9)
ax = axs[1]
for i, (fam, col) in enumerate([('B3L-ligand-1nn', colors[3]), ('B4-concat-mlp', colors[4]), ('dual-encoder', colors[5])]):
    vals = []
    for group in ['new_absent_from_a', 'recurrent']:
        b = res['cells']['declared_increment/screened_primary']['groups'][group]['by_model']
        vals.append(np.mean([v['metrics']['auroc'] for k, v in b.items() if k == fam or k.startswith(fam + '-seed')]))
    ax.plot([0, 1], vals, marker='o', color=col, label=['Target-conditioned 1-NN', 'Concat MLP', 'Dual encoder'][i], lw=1.8)
ax.set(xticks=[0, 1], xticklabels=['Absent from A\n98 ranked targets', 'Recurrent\n29 ranked targets'], ylim=(0.5, 1), ylabel='Macro AUROC', xlim=(-0.18, 1.18))
ax.set_title('b  Different evaluation populations', loc='left', fontsize=11)
ax.legend(fontsize=8, loc='lower right')
ax.grid(axis='y', alpha=0.15)
fig.tight_layout(w_pad=2)
save(fig, 'figure_4_target_and_stratum_diagnostics')
fig, axs = plt.subplots(1, 2, figsize=(9, 3.9), gridspec_kw={'width_ratios': [1.35, 1]})
ax = axs[0]
ax.bar([0, 1, 2], [0.607, 0.624060015327, 0.621], color=['#AABAC7', '#174E77', '#AABAC7'], width=0.55)
ax.errorbar([1], [0.624060015327], yerr=[[0.624060015327 - 0.5781079366], [0.6702409534 - 0.624060015327]], fmt='none', color='black', capsize=5)
ax.axhline(0.7, color='#C76939', ls='--', label='Declared gate 0.70')
ax.set(xticks=[0, 1, 2], xticklabels=['15 Å\nsensitivity', '20 Å\nprimary', '25 Å\nsensitivity'], ylim=(0.5, 0.75), ylabel='Docking AUROC')
ax.legend(loc='upper left', fontsize=9)
ax.set_title('a  Ranking gate failed', loc='left', fontsize=11)
ax = axs[1]
ax.axis('off')
for yy, head, body in [(0.85, 'Pose recovery', '1.514 Å RMSD\nSeparate known-ligand check'), (0.48, 'Scoring coverage', '578 / 600 compounds\n22 failures; selective chemistry'), (0.11, 'Attrition bound', 'Best-case full-cohort AUROC ≤ 0.651\nStill below the 0.70 gate')]:
    ax.text(0.03, yy, head, fontweight='bold', fontsize=11)
    ax.text(0.03, yy - 0.09, body, va='top', fontsize=9.5)
fig.tight_layout()
save(fig, 'figure_5_docking_gate')
provenance['figure_3_family_metrics'] = {label: {'seed_values': v, 'mean': float(np.mean(v))} for label, v in zip(labels, values)}
provenance['figure_4'] = {'rankable_targets': len(x), 'mean_concat': float(np.mean(x)), 'mean_dual': float(np.mean(y)), 'independent_auroc': 'Mann–Whitney rank-sum with half credit for ties'}
provenance['figure_5'] = 'Numbers copied from the accepted docking report; no pose or docking score recomputed here.'
(FIG / 'figure_provenance.json').write_text(json.dumps(provenance, indent=2) + '\n')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.8)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for sn in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[sn]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
st = doc.styles['Normal']
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(7)
st.paragraph_format.line_spacing = 1.08
for sn, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 12), ('Caption', 9)]:
    doc.styles[sn].font.size = Pt(size)
doc.styles['Caption'].paragraph_format.space_after = Pt(9)
header = sec.header.paragraphs[0]
header.text = 'SEQ2LEAD  |  SCIENTIFIC MANUSCRIPT DRAFT'
header.style = 'Caption'
foot = sec.footer.paragraphs[0]
foot.alignment = 2
foot.add_run('Unreviewed draft · 5 October 2026  |  ')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
foot._p.append(fld)
md = []

def p(t, style=None):
    q = doc.add_paragraph(t, style)
    md.append(t + '\n')
    return q

def h(t, l=1):
    doc.add_heading(t, l)
    md.append('#' * (l + 1) + ' ' + t + '\n')

def page():
    doc.add_page_break()

def fig(name, caption):
    q = doc.add_paragraph()
    q.paragraph_format.keep_with_next = True
    q.add_run().add_picture(str(FIG / (name + '.png')), width=Inches(6.85))
    p(caption, 'Caption')
    md.append(f'![{caption}](figures/{name}.svg)\n')

def equation(text):
    q = doc.add_paragraph()
    q.alignment = 1
    om = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = text
    r.append(t)
    om.append(r)
    q._p.append(om)
    md.append('Equation: ' + text + '\n')

def table(headers, data, widths=None):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.autofit = False
    for i, s in enumerate(headers):
        tb.rows[0].cells[i].text = s
    for row in data:
        cells = tb.add_row().cells
        for i, s in enumerate(row):
            cells[i].text = str(s)
    for ri, row in enumerate(tb.rows):
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(4)
                pp.paragraph_format.space_before = Pt(4)
                for rr in pp.runs:
                    rr.font.size = Pt(10)
                    rr.bold = ri == 0
            pr = c._tc.get_or_add_tcPr()
            b = OxmlElement('w:tcBorders')
            for e in ['top', 'left', 'bottom', 'right']:
                z = OxmlElement('w:' + e)
                z.set(qn('w:val'), 'single')
                z.set(qn('w:sz'), '4')
                z.set(qn('w:color'), 'D9D9D9')
                b.append(z)
            pr.append(b)
            if ri == 0:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:fill'), 'EDEFF2')
                pr.append(sh)
        trpr = row._tr.get_or_add_trPr()
        ns = OxmlElement('w:cantSplit')
        trpr.append(ns)
    md.extend(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] + ['| ' + ' | '.join(map(str, r)) + ' |' for r in data] + [''])
p('Seq2Lead historical evaluation of protein ligand ranking', 'Title')
p('Rinat Rizvanov', 'Subtitle')
p('Boston University, Boston, MA, USA', 'Subtitle')
p('Scientific manuscript draft · 5 October 2026 · Not submitted or peer reviewed', 'Caption')
h('Abstract')
p('Sequence-based protein–ligand prediction requires a clear boundary between historical training evidence and later evaluation evidence. We developed Seq2Lead, a provenance-preserving BindingDB pipeline and compound-library ranking tool, and evaluated it in an exploratory historical snapshot study. January and September 2026 snapshots were curated under an identical pinned pipeline. Counted matching distinguished unchanged observations, additions, removals and constrained correction candidates. Training used January evidence alone; increment labels and complete September consistency checks remained separate. Models fitted exact-Ki regression targets and were evaluated with tie-aware per-target ranking metrics. In the primary new-to-fitting cohort, 22,221 pairs covered 992 targets, with 123 meeting the five-positive/five-negative scoring floor. Mean macro AUROC was 0.788873 for a concatenated-feature MLP and 0.781229 for an affine-cosine dual encoder. A separate carbonic-anhydrase-2 docking protocol failed its declared ranking gate despite successful known-pose recovery. The contribution is an auditable implementation and measured case study. Prior corpus exposure and provisional assay pooling preclude a confirmatory interpretation.')
h('Significance')
p('The practical question is whether a screening result can be traced to the evidence, feature content and model selection that produced it. Seq2Lead makes these boundaries inspectable and preserves negative results. It demonstrates that familiar pairs can have substantially different ranking performance from later pairs, and that recovering a crystal pose does not establish useful affinity ranking. These findings support careful evaluation of sequence-based screening tools before experimental prioritisation.')
h('Introduction')
p('Therapeutics Data Commons organises datasets, tasks and benchmarks for therapeutic machine learning [4]. ConPLex established protein-language-model coembedding [1]; Papyrus addresses bioactivity curation [2], and PLINDER addresses protein–ligand evaluation resources [3]. Seq2Lead connects these concerns in a release-aware BindingDB study. It ranks existing compounds from a protein sequence. The aim is to compare a dual encoder with simpler predictors under explicit evidence visibility, while recording how cohort construction and artifact integrity affect interpretation.')
page()
h('Data and evidence design')
fig('figure_1_evidence_workflow', 'Figure 1. Release-aware evidence flow. Historical training, added evaluation evidence and complete later-snapshot consistency are separate readings. Two increment interpretations and two consistency branches are retained; the primary cell is declared increment with screening.')
p('The local September archive loaded 3,237,046 of 3,237,052 data lines, quarantining six with original bytes retained. Curation produced 3,233,963 activity records. Ki, IC50, Kd and EC50 remained distinct. PostgreSQL stored source provenance and curated entities; pinned RDKit processing preserved stated stereochemistry. BindingDB provides the underlying measured affinity evidence [5]. The January deposit (DOI 10.6075/J0V40W61) was ingested and curated in an isolated schema, with read-only exports and source-release/row locators on both sides.')
equation('pKi = 9 − log₁₀(Ki in nM)     ;     active if pKi ≥ 6')
p('Only exact Ki observations supplied regression targets, using the median exact pKi per pair. Censored measurements retained their relation and inclusivity and could decide a class without becoming a numerical target. Empty bound intersections, exact–bound conflicts, ambiguity and discordance were retained in the audit. Exact spread greater than one pKi unit excluded ranking and validation selection under their declared policies; training could retain an exact median. Ki assay poolability was not established, so pooling remained provisional.')
p('The full-B consistency screen reads actual September evidence, including removals and replacements. It does not reconstruct B by appending additions to A. Later corrections and contradictions affect evaluation eligibility and audit records without editing historical training evidence.')
page()
h('Representation and model architecture')
fig('figure_2_dual_encoder', 'Figure 2. Independent compound and protein towers with a trainable affine cosine head. Frozen molecular/protein representations may be computed for every role; learned transforms and model parameters use A-training only. The schematic omits internal layer widths and activation details.')
p('Compounds used chiral ECFP4 fingerprints [6]: 2,048 logical bits stored as 256 packed uint8 values. Proteins used the frozen ESM-2 t33 650M model [7], with mean pooling over residues excluding special tokens and a pinned model commit. Full-length encoding was retained beyond the 1,022-residue training window, flagged as provisional; 40,000 residues was a refusal point. A cache identity combined the computation specification with content manifests. Historical-only feature extensions had separate identities and did not overwrite accepted caches.')
equation('ŷ = a · cos(zc, zp) + b     ;     L = (1/N) Σᵢ (ŷᵢ − pKiᵢ)²')
p('The dual encoder projected both modalities into 512 dimensions and learned scale a and offset b in pKi units. Ranking used the affine score, including the possibility of negative a. Independent towers allow compound projections to be computed before a sequence query. The bounded cosine head limits expressiveness; contrastive training and cross-attention were not implemented. PyTorch neural-network runs used Apple’s MPS GPU backend. CUDA support in the code is not evidence of NVIDIA execution.')
page()
h('Historical matching and fitting protocol')
p('Context slots contained compound, target, normalised publication reference, pH, temperature and source. Canonical measurement values contained type, relation and value. Matching was count-aware: identical rows were indistinguishable observations rather than proof of duplicate experiments. A one-to-one removal/addition could be linked only under the declared identifier rule; ambiguous sets remained unresolved. Entry DOI was entry-level provenance, not a measurement identifier. Equal-valued observations at different slots were correspondence candidates, not proof of one experiment.')
p('Deterministic BLAKE2b slot sharding invoked the same tested matcher for each shard. The recorded full comparison peaked at 0.90 GB RSS, compared with an estimated 25.3 GB unsharded requirement. Pair sharding was used for endpoint aggregation, keeping all evidence for a compound–target pair together. Ki yielded 585,978 unchanged, 33,953 added and 28,777 removed observations. Of the additions, 33,941 entered the declared increment and 12 linked correction after-values were withheld; all stayed in complete B.')
h('Evidence selection algorithm', 2)
q = p('1  Curate and export A and B under the pinned rules.\n2  Match counted observations within context slots.\n3  Derive A train/validation membership without reading B.\n4  Label eligible additions; withhold linked correction occurrences.\n5  Read actual B for the consistency screen and audit.\n6  Derive recurrence from rows supplied to fitting.\n7  Fit on A-train, restore the A-validation-selected model.\n8  Score all declared cells and strata; verify before publishing.')
for rr in q.runs:
    rr.font.name = 'Liberation Mono'
    rr.font.size = Pt(9)
table(['Role or cohort', 'Pairs', 'Purpose'], [['A-train exact regression', '353,957', 'Transform and parameter fitting'], ['A-validation clean regression', '60,981', 'Checkpoint RMSE selection'], ['Primary new to fitting', '22,221', 'Headline ranking evaluation'], ['Strictly absent from A', '19,517', 'Separate exposure stratum']])
p('A deterministic 15% pair reservation supplied validation. The 512-dimensional projection was carried from M9 validation selection with no new sweep. The dual encoder used learning rate 0.001, batch size 512, at most 20 epochs and patience five, restoring the best validation checkpoint. No train-plus-validation refit was performed. Architecture-specific stopping rules were retained; comparisons do not isolate architecture alone. LightGBM [8] runs reached the declared 400-round cap.')
page()
h('Historical ranking results')
fig('figure_3_ranking_results', 'Figure 3. Primary macro AUROC reproduced from saved predictions and labels. Each target contributes equally after meeting the minimum of five active and five inactive pairs. Family means summarise five seeded fits for LightGBM, concat MLP and dual encoder; target mean and target-conditioned 1-NN were fitted once. Seed ranges represent training variation only.')
p('The concatenated-feature MLP reached mean AUROC 0.788873 and the dual encoder 0.781229, an observed difference of 0.007644. The dual encoder did not improve the observed mean. No paired target-level confidence interval, significance test or equivalence test was performed; overlapping seed ranges do not establish equivalence. Both joint configurations exceeded the observed single-modality configurations, but the comparison does not isolate a causal protein-feature effect.')
p('Only 123 of 992 targets cleared the ranking floor; 869 did not. Macro AUROC therefore describes a selected subset with both classes and sufficient observations, not all targets. Pair prevalence was 0.697493, and average target prevalence in the scored subset was 0.543509. Average precision used entire tied threshold blocks, and AUROC gave ties half credit. A protein-only predictor is constant within a target and consequently gives AUROC 0.5 regardless of between-target score differences.')
p('Six families produced 22 predictors. Distinct seeds did not produce identical prediction arrays, even where ranking aggregates were identical. B1 prediction differences were small; B2 predictions differed across seeds while all within-target rankings remained tied. The recorded fits cannot separate stochastic and numerical causes without same-seed repeats. No refits or tuning were undertaken in response to these results.')
page()
h('Target diagnostics and evaluation strata')
fig('figure_4_target_and_stratum_diagnostics', 'Figure 4. (a) Independent rank-sum AUROC calculations for each of the 123 headline targets, averaged across five seeds per family. The dashed line is equality, not a statistical test. (b) Descriptive contrasts between absent-from-A and recurrent populations; the 98 and 29 rankable-target sets differ. Lines connect summaries of different populations rather than longitudinal target trajectories.')
p('The target-conditioned ligand 1-NN searched compounds measured against the same target. It was not a purely ligand-only predictor. Its macro AUROC was 0.964204 on recurrent pairs and 0.668510 on pairs absent from A, a descriptive difference of 0.295694. The dual encoder also performed better on the recurrent population. Chemistry, coverage, label balance and exposure were not independently controlled, so these gaps are not causal estimates of memorisation.')
p('Recurrence was defined against the exact rows emitted to model fitting. Historical presence alone was insufficient: pairs could be reserved for validation, lack regression targets or otherwise be absent from fitting. New-to-fitting includes validation-reserved pairs whose historical evidence participated in model selection. The stricter absent-from-A stratum is therefore reported separately. The validation-reserved subgroup had only six rankable targets and supports little interpretation.')
p('The other three increment/screening cells remain in the saved results. Cross-slot sensitivity removes unambiguous correspondence candidates without calling them the same experiment. These alternatives were declared before eligibility reporting. The paper does not choose an arm retrospectively because it performs better. Feasibility floors are pragmatic requirements, not a statistical power calculation.')
page()
h('Independent docking gate')
fig('figure_5_docking_gate', 'Figure 5. Separate CA2 docking assessment. The 20 Å primary box failed the declared AUROC 0.70 gate. Its bar shows the recorded compound-bootstrap 95% interval; sensitivity bars are rounded summaries and show no inferred interval. The bootstrap does not account for analogue dependence, and its calibration under that dependence is unknown. Pose recovery is a separate check, not a structural image or ranking validation.')
p('AutoDock Vina 1.2.7 [9] used CA2 structure 3K34, whose recorded target sequence matched the curated sequence. A balanced cohort of 600 compounds was frozen before scoring. Of these, 286 actives and 292 inactives scored. Primary AUROC was 0.624060 with recorded interval [0.578108, 0.670241]. Both sensitivity boxes also failed. Even maximising every unknown active–inactive comparison involving attrition gave a full-cohort upper bound of approximately 0.651, below the gate.')
p('Failures were chemically selective. Nineteen unique compounds contained Se, Te or B implicated by the recorded atom-typing errors; per-element counts would double-count two Se/Te compounds. The cohort was overwhelmingly aryl sulfonamides, so generalisation to other chemistry or targets is unsupported. The recorded Mann–Whitney p-value is not a practical effect-size criterion and does not account for analogue dependence.')
p('Known-ligand pose recovery reached 1.514 Å RMSD, independently recalculated from saved pose and crystal evidence in the full local tree. A per-ligand provenance binding from cohort structure to prepared PDBQT was not recorded for existing runs. This assessment does not show whether adding docking to the ML ranking would help; that requires a separate comparison controlling for the model’s existing information.')
page()
h('Discussion and novelty')
p('The contribution is an integrated, auditable implementation and case study: counted historical release changes, three explicit evidence readings, content-bound feature resolution, deterministic memory-bounded processing and publication checks that reject drift. The 28-fold comparison is measured peak RSS against an unsharded estimate, not a benchmark against competing systems. The software also supplies a sequence-query library-ranking demonstration, with model and representation bindings.')
p('The underlying components have established precedents. ConPLex already supports protein-language-model coembedding [1], Papyrus already addresses bioactivity curation [2], and TDC and PLINDER already provide benchmark infrastructure [3,4]. ECFP, ESM-2, LightGBM and Vina are existing methods [6–9]. Seq2Lead introduces no claimed encoder architecture or state-of-the-art method. Its defensible novelty is the specific integration and measured release-aware evaluation, whose broader methodological originality would require further literature review.')
h('Reproducibility and publication integrity', 2)
p('Saved prediction arrays and the consumed evaluation table are checked against fit-recorded digests. The table digest covers labels, strata, eligibility and branch membership, not just pair identity and order. Verification records carry the digests they checked; publication re-hashes current inputs and refuses stale records. Expected input digests remain unchanged, while only derived outputs receive fresh hashes. Fresh scoring must match the published results bytes. Independent rank-sum AUROC provides a second calculation. These controls establish artifact consistency, not independent attestation of the original run.')
h('Limitations and future study', 2)
p('September labels and prior benchmark results were already inspected, and the prospective freeze is unsigned. The historical experiment is exploratory. Poolability across Ki assays, correspondence of cross-slot reports, and long-sequence embedding quality remain unresolved. Retrieval is unbuilt; evaluation evidence display is disabled; the label-reversal split is unscored. Rankability excludes most targets, and no prospective wet-lab validation or external-method comparison was performed. Seed spread is not uncertainty over the target population. A future confirmatory experiment requires a complete dated specification followed by an unseen release or independently withheld labels.')
h('Conclusion', 2)
p('Seq2Lead makes evidence visibility and recorded-artifact consistency inspectable in sequence-based protein–ligand ranking. In the tested historical configurations, the dual encoder did not improve the observed mean over the concat MLP. The separate CA2 docking protocol failed its declared ranking gate despite satisfactory pose recovery. Preserving these bounded findings is a useful result of the evaluation design.')
page()
h('Availability and reproduction')
p('The GitHub repository is pending private creation and owner review. The review package contains code, contracts, saved M11h predictions, evaluation labels and results. It excludes raw source archives, databases, cache vectors, checkpoints and the training membership export. Saved-result recomputation requires no model fit. Checks needing excluded artifacts must be reported unavailable in a fresh clone. The reproduction guide and demo guide distinguish this route from a full local rebuild.')
q = p('uv sync --frozen\nuv run python -m seq2lead.asof.recompute \\\n  data/asof/m11h/run-20261005T134842Z\nuv run python -m seq2lead.asof.verify_results \\\n  data/asof/m11h/run-20261005T134842Z')
for rr in q.runs:
    rr.font.name = 'Liberation Mono'
    rr.font.size = Pt(9)
p('The final contributor list, repository URL and any persistent release identifier remain to be confirmed before circulation. AI coding and review assistants were used extensively during implementation, auditing and drafting. The human author is responsible for the final code, evidence, interpretation and manuscript. This document is a draft technical report, not a submitted or peer-reviewed publication.')
h('References')
refs = [('Singh R, Sledzieski S, Bryson B, Cowen L, Berger B. Contrastive learning in protein language space predicts interactions between drugs and protein targets. PNAS. 2023.', 'https://doi.org/10.1073/pnas.2220778120'), ('Béquignon OJM et al. Papyrus: a large-scale curated dataset aimed at bioactivity predictions. Journal of Cheminformatics. 2023;15:3.', 'https://doi.org/10.1186/s13321-022-00672-x'), ('Durairaj J et al. PLINDER: The protein-ligand interactions dataset and evaluation resource. Preprint. 2024.', 'https://doi.org/10.1101/2024.07.17.603955'), ('Huang K et al. Artificial intelligence foundation for therapeutic science. Nature Chemical Biology. 2022;18:1033–1036.', 'https://doi.org/10.1038/s41589-022-01131-2'), ('BindingDB in 2024: a FAIR knowledgebase of protein-small molecule binding data. Nucleic Acids Research. 2025;53(D1):D1633–D1644.', 'https://doi.org/10.1093/nar/gkae1075'), ('Rogers D, Hahn M. Extended-connectivity fingerprints. Journal of Chemical Information and Modeling. 2010;50:742–754.', 'https://doi.org/10.1021/ci100050t'), ('Lin Z et al. Evolutionary-scale prediction of atomic-level protein structure with a language model. Science. 2023;379:1123–1130.', 'https://doi.org/10.1126/science.ade2574'), ('Ke G et al. LightGBM: A Highly Efficient Gradient Boosting Decision Tree. Advances in Neural Information Processing Systems. 2017;30.', 'https://proceedings.neurips.cc/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html'), ('Eberhardt J et al. AutoDock Vina 1.2.0: New Docking Methods, Expanded Force Field, and Python Bindings. Journal of Chemical Information and Modeling. 2021;61:3891–3898.', 'https://doi.org/10.1021/acs.jcim.1c00203')]
for i, (t, url) in enumerate(refs, 1):
    q = p(f'{i}. {t}', 'Normal')
    q.paragraph_format.space_after = Pt(4)
    for rr in q.runs:
        rr.font.size = Pt(9)
    link = OxmlElement('w:hyperlink')
    rid = q.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link.set(qn('r:id'), rid)
    rr = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '18')
    rp.append(sz)
    rr.append(rp)
    tt = OxmlElement('w:t')
    tt.text = ' ' + url
    tt.set(qn('xml:space'), 'preserve')
    rr.append(tt)
    link.append(rr)
    q._p.append(link)
    md[-1] = f'{i}. {t} {url}\n'
doc.core_properties.title = 'Seq2Lead historical evaluation of protein ligand ranking'
doc.core_properties.author = 'Rinat Rizvanov'
doc.core_properties.subject = 'Exploratory historical BindingDB evaluation and artifact integrity'
out = ROOT / 'Seq2Lead_scientific_manuscript.docx'
doc.save(out)
(REPO / 'paper/manuscript.md').write_text('\n'.join(md), encoding='utf-8')
(REPO / 'paper/FIGURES.md').write_text('# Figure guide\n\nFigures 1 and 2 are original explanatory schematics. Figures 3 and 4 use digest-verified saved M11h predictions and evaluation labels; the target AUROC calculation independently reproduces the family means to 1e-12. Figure 5 plots accepted reported M10 values without rerunning docking. All figures have PNG and editable SVG versions; text and shapes in SVG can be edited with a vector editor. The manuscript text, tables and equations are editable in Word; its figures are embedded images.\n\n`figure_provenance.json` records inputs and checks; `per_target_comparison.json` exposes the plotted target summaries. Run `python paper/build_paper.py` from the repository using Python with numpy, scipy, matplotlib and python-docx installed. The script uses only saved artifacts and fits no models.\n\nNot constructed: predicted-versus-measured pKi/R² (numerical evaluation targets absent from this package), probability calibration (outputs are pKi, not calibrated probabilities), embedding maps (vectors absent), atom attribution (not computed), and pose overlays (coordinates absent). These require the corresponding original artifacts or a separately scoped analysis. No synthetic figure represents an observed molecule, embedding or pose.\n\nTDC figure design was consulted through the public paper *Artificial intelligence foundation for therapeutic science* (2022), figures 1–2. No published artwork was copied. The earlier uploaded papers were not recoverable in this session.\n')
print(out)
print('independent AUROC', len(x), np.mean(x), np.mean(y))
